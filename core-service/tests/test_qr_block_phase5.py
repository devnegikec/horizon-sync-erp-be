"""Phase 5 tests for external QR short URL generation."""

import json
from types import SimpleNamespace
from unittest.mock import Mock
from uuid import uuid4

import httpx
import pytest

from app.services.qr_product_service import QRProductService
from app.services.qr_shortener import QRShortener, QRShortenerError


def _client(handler) -> httpx.Client:
    return httpx.Client(transport=httpx.MockTransport(handler))


def _service() -> QRProductService:
    service = QRProductService.__new__(QRProductService)
    service.db = Mock()
    service.item_repo = Mock()
    service.key_service = Mock()
    service.key_service.decrypt_private_key.return_value = object()
    service.key_service.sign_message.return_value = "signed-value"
    service.qr_shortener = Mock()
    return service


def _block(qr_type: str = "dynamic") -> SimpleNamespace:
    return SimpleNamespace(
        id=uuid4(),
        product_id=uuid4(),
        sku_id=None,
        quantity=1,
        batch="PHASE5",
        qr_type=qr_type,
        sr_number_type="R8DAN",
        starting_serial=None,
        serial_prefix="PRO",
    )


def _product() -> SimpleNamespace:
    return SimpleNamespace(
        brand_id=uuid4(),
        gtin="0123456789012",
        sr_number_type="R8DAN",
        activation_method="pre",
    )


def _brand(monkeypatch) -> None:
    monkeypatch.setattr(
        "app.repositories.brand_repository.BrandRepository.get_by_id",
        Mock(
            return_value=SimpleNamespace(
                private_key_encrypted="encrypted",
                short_code="demo",
            )
        ),
    )


def test_shortener_posts_existing_cloudfront_contract(monkeypatch):
    received = {}

    def handler(request: httpx.Request) -> httpx.Response:
        received["url"] = str(request.url)
        received["payload"] = json.loads(request.read())
        return httpx.Response(200, json={"url_short": "https://bwqr.me/abc"})

    monkeypatch.setattr("app.services.qr_shortener.settings.qr_shortener_enabled", True)
    short_url = QRShortener(client=_client(handler)).shorten(
        "https://demo.verify.example.com/g/1/s/ABC/123?c=signature"
    )

    assert short_url == "https://bwqr.me/abc"
    assert received["url"] == "https://de5be4rdmboho.cloudfront.net/prod/"
    assert received["payload"] == {
        "cdn_prefix": "bwqr.me",
        "url_long": ("https://demo.verify.example.com/g/1/s/ABC/123?c=signature"),
    }


def test_shortener_retries_transient_failure(monkeypatch):
    responses = iter(
        [
            httpx.Response(503),
            httpx.Response(200, json={"url_short": "https://bwqr.me/recovered"}),
        ]
    )
    sleeps = []
    monkeypatch.setattr("app.services.qr_shortener.settings.qr_shortener_enabled", True)

    short_url = QRShortener(
        client=_client(lambda _request: next(responses)),
        sleep=sleeps.append,
    ).shorten("https://demo.verify.example.com/long")

    assert short_url == "https://bwqr.me/recovered"
    assert sleeps == [0.25]


def test_shortener_normalizes_provider_http_url_to_https(monkeypatch):
    monkeypatch.setattr("app.services.qr_shortener.settings.qr_shortener_enabled", True)
    client = _client(
        lambda _request: httpx.Response(
            200,
            json={"url_short": "http://bwqr.me/provider-code"},
        )
    )

    short_url = QRShortener(client=client).shorten(
        "https://demo.verify.example.com/long"
    )

    assert short_url == "https://bwqr.me/provider-code"


@pytest.mark.parametrize(
    "response",
    [
        httpx.Response(200, json={"unexpected": "value"}),
        httpx.Response(200, json={"url_short": "ftp://bwqr.me/invalid"}),
        httpx.Response(200, json={"url_short": "https://other.test/item"}),
        httpx.Response(400, json={"message": "bad request"}),
    ],
)
def test_shortener_rejects_invalid_or_non_retryable_responses(response, monkeypatch):
    calls = 0

    def handler(_request: httpx.Request) -> httpx.Response:
        nonlocal calls
        calls += 1
        return response

    monkeypatch.setattr("app.services.qr_shortener.settings.qr_shortener_enabled", True)
    with pytest.raises(QRShortenerError, match="QR short URL generation failed"):
        QRShortener(client=_client(handler), sleep=lambda _delay: None).shorten(
            "https://demo.verify.example.com/long"
        )

    assert calls == 1


def test_generated_item_uses_short_url_and_retains_long_url(monkeypatch):
    service = _service()
    _brand(monkeypatch)
    service.item_repo.get_existing_serials_global.return_value = set()
    service.qr_shortener.shorten.return_value = "https://bwqr.me/item"

    service._generate_product_items(_block(), _product(), uuid4(), uuid4())

    item = service.item_repo.bulk_create.call_args.args[0][0]
    long_url = item["extra_data"]["long_url"]
    assert item["token_id"] == "https://bwqr.me/item"
    assert item["extra_data"]["short_url"] == item["token_id"]
    assert "/g/0123456789012/s/" in long_url
    service.qr_shortener.shorten.assert_called_once_with(long_url)


def test_dual_item_shortens_both_urls_and_retains_both_long_urls(monkeypatch):
    service = _service()
    _brand(monkeypatch)
    service.item_repo.get_existing_serials_global.return_value = set()
    service.qr_shortener.shorten.side_effect = [
        "https://bwqr.me/overt",
        "https://bwqr.me/covert",
    ]

    service._generate_product_items(_block("dual"), _product(), uuid4(), uuid4())

    item = service.item_repo.bulk_create.call_args.args[0][0]
    metadata = item["extra_data"]
    assert item["token_id"] == "https://bwqr.me/overt"
    assert metadata["overt_url"] == "https://bwqr.me/overt"
    assert metadata["covert_url"] == "https://bwqr.me/covert"
    assert metadata["overt_long_url"].endswith("&qr=overt")
    assert metadata["covert_long_url"].endswith("&qr=covert")
    assert "/g/0123456789012/s/" in metadata["covert_long_url"]
    assert service.qr_shortener.shorten.call_count == 2


@pytest.mark.parametrize(
    "qr_type",
    ["dynamic", "static", "secure_code", "dual", "one_time", "post_activation"],
)
def test_redirect_mode_keeps_serial_bearing_qr_url(monkeypatch, qr_type):
    """Redirect products still encode the /g/{gtin}/s/{serial} payload.

    Direct redirect happens when the QR is scanned, so the QR structure must
    stay identical to the pre-existing (dev) format for every QR type.
    """
    service = _service()
    _brand(monkeypatch)
    service.item_repo.get_existing_serials_global.return_value = set()
    service.qr_shortener.shorten.return_value = "https://bwqr.me/item"
    product = _product()
    product.redirect_to_client = True
    product.client_product_auth_url = "https://client.example.com/product"

    service._generate_product_items(_block(qr_type), product, uuid4(), uuid4())

    item = service.item_repo.bulk_create.call_args.args[0][0]
    metadata = item["extra_data"]
    assert "/g/0123456789012/s/" in metadata["long_url"]
    assert "client.example.com" not in json.dumps(metadata)
    if qr_type == "dual":
        assert "/g/0123456789012/s/" in metadata["covert_long_url"]
        assert metadata["overt_url"] == "https://bwqr.me/item"
    else:
        assert item["token_id"] == "https://bwqr.me/item"
    service.key_service.decrypt_private_key.assert_called_once()


def test_shortener_failure_fails_block_and_releases_reserved_credits():
    service = QRProductService.__new__(QRProductService)
    service.db = Mock()
    service.block_repo = Mock()
    service.item_repo = Mock()
    service.product_repo = Mock()
    service.credit_service = Mock()
    block = SimpleNamespace(
        id=uuid4(),
        organization_id=uuid4(),
        product_id=uuid4(),
        created_by=uuid4(),
        status="pending",
        task_status="pending",
        task_id="phase5-task",
        quantity=1,
        progress=0,
        generated_count=0,
        completed_at=None,
        error_code=None,
        error_message=None,
        artifact_object_key=None,
        artifact_size_bytes=None,
        artifact_checksum_sha256=None,
        artifact_generated_at=None,
    )
    service.block_repo.get_by_id_for_update.return_value = block
    service.block_repo.get_by_id.return_value = block
    service.product_repo.get_by_id.return_value = SimpleNamespace(id=block.product_id)
    service._generate_product_items = Mock(
        side_effect=QRShortenerError("QR short URL generation failed")
    )

    with pytest.raises(QRShortenerError):
        service.process_block(
            block.id,
            block.organization_id,
            task_id=block.task_id,
        )

    assert block.status == "failed"
    assert block.error_code == "generation_failed"
    service.credit_service.release_reserved_credits.assert_called_once_with(
        block.organization_id,
        block.id,
    )
    service.credit_service.consume_reserved_credits.assert_not_called()


class TestQrPayloadUsesSerialBearingUrl:
    """The QR image must encode a URL the handheld app can parse.

    Regression guard: the CDN short link (``/01/{gtin}/{slug}``) was being written
    into the QR, but neither the app nor QSeal parent→child linking can read it —
    both expect ``/g/{gtin}/s/{serial}``.
    """

    SHORT = "https://bwqr.me/01/8843266542/ZQ21nGQDYb4aId5o"
    LONG = "https://horizon.ciphercode.ai/g/8843266542/s/PHF-SFOTO8/1792048740930?c=sig"

    def test_branded_item_encodes_long_url_not_short_url(self):
        rows = QRProductService._items_to_excel_rows(
            [
                {
                    "serial_number": "PHF-SFOTO8",
                    "token_id": self.SHORT,
                    "extra_data": {"long_url": self.LONG, "short_url": self.SHORT},
                }
            ]
        )

        assert len(rows) == 1
        assert rows[0]["primary_url"] == self.LONG
        assert "/g/8843266542/s/PHF-SFOTO8" in rows[0]["primary_url"]
        assert rows[0]["primary_url"] != self.SHORT

    def test_unsigned_block_falls_back_to_token(self):
        """No brand key -> no long_url; token_id is already serial-bearing."""
        url = "https://horizon.ciphercode.ai/g/0123456789012/s/PRO-ABC123"
        rows = QRProductService._items_to_excel_rows(
            [
                {
                    "serial_number": "PRO-ABC123",
                    "token_id": url,
                    "extra_data": None,
                }
            ]
        )

        assert rows[0]["primary_url"] == url

    def test_item_without_any_url_stays_empty(self):
        rows = QRProductService._items_to_excel_rows(
            [{"serial_number": "X", "token_id": None, "extra_data": {}}]
        )
        assert rows[0]["primary_url"] == ""
        assert rows[0]["covert_url"] == ""

    def test_dual_item_encodes_both_long_urls(self):
        short_overt = "https://bwqr.me/01/1/overt"
        short_covert = "https://bwqr.me/01/1/covert"
        long_overt = "https://horizon.ciphercode.ai/g/1/s/SER/1?c=s&qr=overt"
        long_covert = "https://horizon.ciphercode.ai/g/1/s/SER/1?c=s&qr=covert"

        rows = QRProductService._items_to_excel_rows(
            [
                {
                    "serial_number": "SER",
                    "token_id": short_overt,
                    "extra_data": {
                        "overt_url": short_overt,
                        "covert_url": short_covert,
                        "overt_long_url": long_overt,
                        "covert_long_url": long_covert,
                    },
                }
            ]
        )

        assert rows[0]["overt_url"] == long_overt
        assert rows[0]["covert_url"] == long_covert
        assert rows[0]["overt_url"] != rows[0]["covert_url"]

    def test_orm_backed_item_is_handled(self):
        """Download path passes ProductItem ORM rows, not dicts."""
        item = SimpleNamespace(
            serial_number="PHF-SFOTO8",
            token_id=self.SHORT,
            secrete_code=None,
            extra_data={"long_url": self.LONG, "short_url": self.SHORT},
        )

        rows = QRProductService._items_to_excel_rows([item])

        assert rows[0]["serial"] == "PHF-SFOTO8"
        assert rows[0]["primary_url"] == self.LONG

    def test_encoded_row_builds_workbook_with_long_url(self):
        """End-to-end: the long URL lands in the workbook that renders the QR."""
        from io import BytesIO

        from openpyxl import load_workbook

        from app.services.qr_product_service import _build_excel

        rows = QRProductService._items_to_excel_rows(
            [
                {
                    "serial_number": "PHF-SFOTO8",
                    "token_id": self.SHORT,
                    "extra_data": {"long_url": self.LONG, "short_url": self.SHORT},
                }
            ]
        )
        data = _build_excel(rows, "dynamic", include_qr_images=True)
        worksheet = load_workbook(BytesIO(data)).active

        assert worksheet.cell(2, 1).value == self.LONG
        assert worksheet.cell(2, 2).value == "PHF-SFOTO8"
        assert len(worksheet._images) == 1

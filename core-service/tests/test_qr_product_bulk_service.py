"""Unit tests for QR product SKU normalisation and bulk import/export parsing.

Only the database-free logic is covered here: SKU normalisation, per-row
parsing/validation, and CSV/XLSX rendering + reading. The upsert path itself is
exercised through ``QRProductService`` against a real database.
"""

import csv
import io
from decimal import Decimal
from uuid import uuid4

import pytest

from app.services.qr_product_bulk_service import (
    PRODUCT_COLUMNS,
    SAMPLE_ROWS,
    QRProductBulkService,
    parse_row,
)
from app.services.qr_product_service import normalize_sku


class TestNormalizeSku:
    def test_trims_and_treats_blank_as_unset(self):
        assert normalize_sku("  ABC-1  ") == "ABC-1"
        assert normalize_sku("   ") is None
        assert normalize_sku("") is None
        assert normalize_sku(None) is None


class TestParseRow:
    def test_requires_name(self):
        with pytest.raises(ValueError, match="name is required"):
            parse_row({"sku": "ABC-1"}, 2)

    def test_omits_blank_fields_and_normalises_sku(self):
        payload = parse_row({"name": " Widget ", "sku": "  abc-1 "}, 2)
        assert payload["name"] == "Widget"
        assert payload["sku"] == "abc-1"
        # Unprovided columns are absent (so an upsert never blanks them).
        assert "generic_name" not in payload

    def test_rejects_invalid_activation_method(self):
        with pytest.raises(ValueError, match="activation_method"):
            parse_row({"name": "W", "activation_method": "sometimes"}, 2)

    def test_rejects_invalid_brand_id(self):
        with pytest.raises(ValueError, match="brand_id"):
            parse_row({"name": "W", "brand_id": "not-a-uuid"}, 2)

    def test_keeps_a_valid_brand_id(self):
        brand_id = uuid4()
        payload = parse_row({"name": "W", "brand_id": str(brand_id)}, 2)
        assert payload["brand_id"] == brand_id

    def test_rejects_fractional_integer_field(self):
        # ``int(Decimal("1.9"))`` used to truncate to 1 silently.
        with pytest.raises(ValueError, match="items_per_master_pack"):
            parse_row({"name": "W", "items_per_master_pack": "1.9"}, 2)
        with pytest.raises(ValueError, match="warranty_period_months"):
            parse_row({"name": "W", "warranty_period_months": "2.5"}, 2)

    def test_settings_accept_id_or_value(self):
        setting_id = uuid4()
        by_id = parse_row({"name": "W", "shelf_life_setting_id": str(setting_id)}, 2)
        assert by_id["shelf_life_setting_id"] == setting_id

        by_value = parse_row({"name": "W", "serial_prefix_setting_value": "PH"}, 2)
        assert by_value["serial_prefix_setting_value"] == "PH"

        # A code placed in the *_id column is tolerated as a value.
        tolerated = parse_row({"name": "W", "shelf_life_setting_id": "12"}, 2)
        assert tolerated["shelf_life_setting_value"] == "12"

    def test_packaging_defaults_when_any_field_present(self):
        payload = parse_row({"name": "W", "items_per_master_pack": "8"}, 2)
        assert payload["packaging"] == {
            "items_per_master_pack": 8,
            "unit_name": "Each",
            "conversion_factor": Decimal("1"),
        }

    def test_packaging_absent_when_no_fields(self):
        assert "packaging" not in parse_row({"name": "W"}, 2)

    def test_rejects_non_positive_conversion_factor(self):
        with pytest.raises(ValueError, match="conversion_factor"):
            parse_row({"name": "W", "conversion_factor": "0"}, 2)

    # ── Master carton (MC) ────────────────────────────────────────────────

    def test_master_carton_fields_are_parsed(self):
        payload = parse_row(
            {
                "name": "W",
                "conversion_factor": "12",
                "master_pack_unit_name": "Carton of 12",
                "master_pack_length_mm": "400",
                "master_pack_width_mm": "300",
                "master_pack_height_mm": "250",
                "master_pack_weight_grams": "5400",
                "master_pack_fill_factor": "0.8",
                "master_pack_void_fill_pct": "0.15",
                "master_pack_wall_thickness_mm": "4",
            },
            2,
        )
        packaging = payload["packaging"]
        assert packaging["master_pack_unit_name"] == "Carton of 12"
        assert packaging["master_pack_length_mm"] == Decimal("400")
        assert packaging["master_pack_width_mm"] == Decimal("300")
        assert packaging["master_pack_height_mm"] == Decimal("250")
        assert packaging["master_pack_weight_grams"] == Decimal("5400")
        assert packaging["master_pack_fill_factor"] == Decimal("0.8")
        assert packaging["master_pack_void_fill_pct"] == Decimal("0.15")
        assert packaging["master_pack_wall_thickness_mm"] == Decimal("4")
        # IC defaults are still applied whenever any packaging column is present.
        assert packaging["unit_name"] == "Each"
        assert packaging["conversion_factor"] == Decimal("12")

    def test_master_carton_only_row_requires_a_pack_size(self):
        # Without items_per_master_pack (or conversion_factor > 1) the item
        # service deactivates the carton and silently drops the MC data, so the
        # row must be rejected instead of accepted and discarded.
        with pytest.raises(ValueError, match="items_per_master_pack"):
            parse_row({"name": "W", "master_pack_unit_name": "Carton"}, 2)
        with pytest.raises(ValueError, match="items_per_master_pack"):
            parse_row(
                {"name": "W", "master_pack_length_mm": "400", "conversion_factor": "1"},
                2,
            )

    def test_master_carton_only_row_with_pack_size_is_accepted(self):
        payload = parse_row(
            {
                "name": "W",
                "master_pack_unit_name": "Carton",
                "items_per_master_pack": "12",
            },
            2,
        )
        assert payload["packaging"]["master_pack_unit_name"] == "Carton"
        assert payload["packaging"]["items_per_master_pack"] == 12
        assert payload["packaging"]["unit_name"] == "Each"

    def test_master_carton_row_accepts_conversion_factor_above_one(self):
        payload = parse_row(
            {
                "name": "W",
                "master_pack_unit_name": "Carton",
                "conversion_factor": "12",
            },
            2,
        )
        assert payload["packaging"]["conversion_factor"] == Decimal("12")

    def test_rejects_master_pack_fill_factor_out_of_range(self):
        with pytest.raises(ValueError, match="master_pack_fill_factor"):
            parse_row({"name": "W", "master_pack_fill_factor": "0"}, 2)
        with pytest.raises(ValueError, match="master_pack_fill_factor"):
            parse_row({"name": "W", "master_pack_fill_factor": "1.2"}, 2)

    def test_rejects_master_pack_void_fill_out_of_range(self):
        with pytest.raises(ValueError, match="master_pack_void_fill_pct"):
            parse_row({"name": "W", "master_pack_void_fill_pct": "1.5"}, 2)

    def test_rejects_negative_master_pack_dimension(self):
        with pytest.raises(ValueError, match="master_pack_length_mm"):
            parse_row({"name": "W", "master_pack_length_mm": "-1"}, 2)

    # ── Linked inventory item fields ──────────────────────────────────────

    def test_item_fields_are_parsed(self):
        group_id = uuid4()
        payload = parse_row(
            {
                "name": "W",
                "description": "A widget",
                "uom": "Box",
                "item_group_id": str(group_id),
                "maintain_stock": "false",
                "valuation_method": "FIFO",
                "standard_rate": "19.99",
                "valuation_rate": "12.50",
                "min_order_qty": "5",
                "max_order_qty": "500",
                "reorder_level": "20",
                "reorder_qty": "100",
                "weight_per_unit": "0.75",
                "weight_uom": "kg",
                "barcode": "8901234567890",
                "image_url": "https://example.com/a.jpg",
                "item_status": "Active",
            },
            2,
        )
        fields = payload["item_fields"]
        assert fields["description"] == "A widget"
        assert fields["uom"] == "Box"
        assert fields["item_group_id"] == group_id
        assert fields["maintain_stock"] is False
        assert fields["valuation_method"] == "fifo"
        assert fields["standard_rate"] == Decimal("19.99")
        assert fields["valuation_rate"] == Decimal("12.50")
        assert fields["min_order_qty"] == 5
        assert fields["max_order_qty"] == 500
        assert fields["reorder_level"] == 20
        assert fields["reorder_qty"] == 100
        assert fields["weight_per_unit"] == Decimal("0.75")
        assert fields["weight_uom"] == "kg"
        assert fields["barcode"] == "8901234567890"
        assert fields["image_url"] == "https://example.com/a.jpg"
        assert fields["status"] == "active"

    def test_item_group_name_is_kept_for_later_resolution(self):
        payload = parse_row({"name": "W", "item_group_name": "Electronics"}, 2)
        assert payload["item_fields"]["item_group_name"] == "Electronics"

    def test_item_fields_absent_when_no_item_columns(self):
        assert "item_fields" not in parse_row({"name": "W", "sku": "ABC-1"}, 2)

    def test_rejects_invalid_item_group_id(self):
        with pytest.raises(ValueError, match="item_group_id"):
            parse_row({"name": "W", "item_group_id": "not-a-uuid"}, 2)

    def test_rejects_invalid_valuation_method(self):
        with pytest.raises(ValueError, match="valuation_method"):
            parse_row({"name": "W", "valuation_method": "guesswork"}, 2)

    def test_rejects_invalid_item_status(self):
        with pytest.raises(ValueError, match="item_status"):
            parse_row({"name": "W", "item_status": "archived"}, 2)

    def test_rejects_negative_rate(self):
        with pytest.raises(ValueError, match="standard_rate"):
            parse_row({"name": "W", "standard_rate": "-1"}, 2)

    def test_rejects_min_order_qty_below_one(self):
        with pytest.raises(ValueError, match="min_order_qty"):
            parse_row({"name": "W", "min_order_qty": "0"}, 2)


class TestRenderAndRead:
    def _service(self):
        # db is unused by the pure rendering/reading helpers.
        return QRProductBulkService(db=None)

    def test_csv_export_has_header_and_rows(self):
        svc = self._service()
        content, filename = svc._render(
            [{"name": "Widget", "sku": "ABC-1"}], PRODUCT_COLUMNS, "csv", "exp"
        )
        assert filename.endswith(".csv")
        text = content.decode("utf-8-sig")
        header = text.splitlines()[0]
        assert header.split(",")[0] == "name"
        assert "Widget" in text
        assert "ABC-1" in text

    def test_xlsx_export_round_trips_headers(self):
        from openpyxl import load_workbook

        svc = self._service()
        content, filename = svc._render(
            [{"name": "Widget", "sku": "ABC-1"}], PRODUCT_COLUMNS, "xlsx", "exp"
        )
        assert filename.endswith(".xlsx")
        wb = load_workbook(io.BytesIO(content), read_only=True)
        rows = list(wb.active.iter_rows(values_only=True))
        assert rows[0] == tuple(PRODUCT_COLUMNS)
        assert rows[1][0] == "Widget"

    def test_read_rows_csv_skips_empty_rows(self):
        svc = self._service()
        csv_bytes = (
            b"name,sku\n"
            b"Widget,ABC-1\n"
            b",\n"  # blank row must be ignored
            b"Gadget,ABC-2\n"
        )
        rows = svc._read_rows(csv_bytes, "products.csv")
        assert [row["name"] for row in rows] == ["Widget", "Gadget"]

    def test_read_rows_xlsx_skips_empty_rows(self):
        svc = self._service()
        content, _ = svc._render(
            [{"name": "Widget", "sku": "ABC-1"}], PRODUCT_COLUMNS, "xlsx", "t"
        )
        rows = svc._read_rows(content, "products.xlsx")
        assert len(rows) == 1 and rows[0]["name"] == "Widget"

    def test_read_rows_rejects_unknown_extension(self):
        with pytest.raises(ValueError, match="Unsupported file type"):
            self._service()._read_rows(b"data", "products.pdf")

    def test_read_rows_rejects_empty_file(self):
        with pytest.raises(ValueError, match="empty"):
            self._service()._read_rows(b"", "products.csv")

    def test_create_kwargs_include_brand_id(self):
        brand_id = uuid4()
        kwargs = QRProductBulkService._create_kwargs(
            {"name": "W", "sku": "ABC-1", "brand_id": brand_id}
        )
        assert kwargs["brand_id"] == brand_id

    def test_export_neutralises_formula_injection(self):
        svc = self._service()
        content, _ = svc._render(
            [{"name": '=HYPERLINK("http://evil")', "sku": "+cmd|calc"}],
            PRODUCT_COLUMNS,
            "csv",
            "exp",
        )
        text = content.decode("utf-8-sig")
        assert "'=HYPERLINK" in text
        assert "'+cmd|calc" in text

    def test_export_leaves_ordinary_text_untouched(self):
        svc = self._service()
        content, _ = svc._render(
            [{"name": "Widget", "sku": "ABC-1"}], PRODUCT_COLUMNS, "csv", "exp"
        )
        assert "Widget" in content.decode("utf-8-sig")


class TestTemplateSamples:
    """The template ships example rows, so guard that they stay importable."""

    def _template_rows(self, file_format: str = "csv") -> list[dict]:
        svc = QRProductBulkService(db=None)
        content, _ = svc.template(file_format)
        text = content.decode("utf-8-sig")
        return list(csv.DictReader(io.StringIO(text)))

    def test_template_ships_header_and_samples(self):
        rows = self._template_rows()
        assert len(rows) == len(SAMPLE_ROWS)
        assert list(rows[0].keys()) == PRODUCT_COLUMNS

    def test_xlsx_template_ships_samples(self):
        from openpyxl import load_workbook

        svc = QRProductBulkService(db=None)
        content, _ = svc.template("xlsx")
        wb = load_workbook(io.BytesIO(content), read_only=True)
        sheet_rows = list(wb.active.iter_rows(values_only=True))
        assert sheet_rows[0] == tuple(PRODUCT_COLUMNS)
        assert len(sheet_rows) == len(SAMPLE_ROWS) + 1

    def test_every_sample_row_parses(self):
        for index, row in enumerate(self._template_rows(), start=2):
            payload = parse_row(row, index)
            assert payload["name"]
            assert payload.get("sku")

    def test_samples_cover_ic_mc_and_item_columns(self):
        packed = SAMPLE_ROWS[0]
        assert packed["items_per_master_pack"]
        assert packed["master_pack_unit_name"]
        assert packed["standard_rate"]


class TestPackageColumns:
    def test_template_columns_cover_ic_mc_and_item_fields(self):
        expected = [
            # Inner carton (IC)
            "unit_name",
            "conversion_factor",
            "items_per_master_pack",
            "length_mm",
            "width_mm",
            "height_mm",
            "weight_grams",
            # Master carton (MC)
            "master_pack_unit_name",
            "master_pack_length_mm",
            "master_pack_width_mm",
            "master_pack_height_mm",
            "master_pack_weight_grams",
            "master_pack_fill_factor",
            "master_pack_void_fill_pct",
            "master_pack_wall_thickness_mm",
            # Linked item
            "description",
            "uom",
            "item_group_id",
            "item_group_name",
            "maintain_stock",
            "valuation_method",
            "standard_rate",
            "valuation_rate",
            "min_order_qty",
            "max_order_qty",
            "reorder_level",
            "reorder_qty",
            "weight_per_unit",
            "weight_uom",
            "barcode",
            "item_status",
        ]
        for column in expected:
            assert column in PRODUCT_COLUMNS

    def test_columns_are_unique(self):
        assert len(PRODUCT_COLUMNS) == len(set(PRODUCT_COLUMNS))


class TestPackagingFromItem:
    class _Unit:
        def __init__(self, **kwargs):
            self.unit_name = "Each"
            self.conversion_factor = Decimal("1")
            self.items_per_master_pack = None
            self.length_mm = None
            self.width_mm = None
            self.height_mm = None
            self.weight_grams = None
            self.master_pack_fill_factor = None
            self.master_pack_void_fill_pct = None
            self.master_pack_wall_thickness_mm = None
            self.is_base_unit = False
            self.is_active = True
            for key, value in kwargs.items():
                setattr(self, key, value)

    def test_splits_base_unit_and_carton(self):
        from app.services.qr_product_bulk_service import _packaging_from_item

        base = self._Unit(
            unit_name="Each",
            conversion_factor=Decimal("1"),
            items_per_master_pack=12,
            length_mm=Decimal("100"),
            is_base_unit=True,
        )
        carton = self._Unit(
            unit_name="Carton of 12",
            conversion_factor=Decimal("12"),
            length_mm=Decimal("400"),
            weight_grams=Decimal("5000"),
            master_pack_fill_factor=Decimal("0.8"),
            is_base_unit=False,
        )
        item = type("Item", (), {"packaging_units": [base, carton]})()

        ic, mc = _packaging_from_item(item)
        assert ic["unit_name"] == "Each"
        assert ic["items_per_master_pack"] == 12
        assert mc["master_pack_unit_name"] == "Carton of 12"
        assert mc["master_pack_length_mm"] == Decimal("400")
        assert mc["master_pack_weight_grams"] == Decimal("5000")
        assert mc["master_pack_fill_factor"] == Decimal("0.8")
        # The MC dict must not leak the un-prefixed IC keys.
        assert "unit_name" not in mc

    def test_ignores_inactive_units(self):
        from app.services.qr_product_bulk_service import _packaging_from_item

        inactive = self._Unit(unit_name="Old Carton", is_active=False)
        item = type("Item", (), {"packaging_units": [inactive]})()
        ic, mc = _packaging_from_item(item)
        assert ic == {} and mc == {}

    def test_picks_the_outermost_carton_deterministically(self):
        from app.services.qr_product_bulk_service import _packaging_from_item

        box = self._Unit(unit_name="Box of 12", conversion_factor=Decimal("12"))
        pallet = self._Unit(unit_name="Pallet of 144", conversion_factor=Decimal("144"))
        # Order in the relationship is not guaranteed; the outer unit must win
        # regardless of the order the rows come back in.
        for units in ([box, pallet], [pallet, box]):
            item = type("Item", (), {"packaging_units": units})()
            _, mc = _packaging_from_item(item)
            assert mc["master_pack_unit_name"] == "Pallet of 144"

    def test_returns_empty_dicts_without_item(self):
        from app.services.qr_product_bulk_service import _packaging_from_item

        assert _packaging_from_item(None) == ({}, {})


class TestItemPackagingMapping:
    def test_unset_master_pack_knobs_are_not_forwarded(self):
        # None must stay None on the item schema, otherwise every partial
        # packaging update would reset stored MC settings to schema defaults.
        from app.services.qr_product_service import _to_item_packaging_details

        details = _to_item_packaging_details(
            {
                "unit_name": "Each",
                "conversion_factor": Decimal("1"),
                "master_pack_fill_factor": None,
                "master_pack_void_fill_pct": None,
                "master_pack_wall_thickness_mm": None,
            }
        )
        assert details.master_pack_fill_factor is None
        assert details.master_pack_void_fill_pct is None
        assert details.master_pack_wall_thickness_mm is None

    def test_explicit_master_pack_knobs_are_forwarded(self):
        from app.services.qr_product_service import _to_item_packaging_details

        details = _to_item_packaging_details(
            {
                "unit_name": "Each",
                "conversion_factor": Decimal("1"),
                "master_pack_fill_factor": Decimal("0.8"),
                "master_pack_void_fill_pct": Decimal("0.15"),
                "master_pack_wall_thickness_mm": Decimal("4"),
            }
        )
        assert details.master_pack_fill_factor == Decimal("0.8")
        assert details.master_pack_void_fill_pct == Decimal("0.15")
        assert details.master_pack_wall_thickness_mm == Decimal("4")

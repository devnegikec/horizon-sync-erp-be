"""Bulk import / export for QR products (CSV + XLSX).

Design notes
------------
* **SKU is the upsert key.** A row whose (case/whitespace-insensitive) SKU
  matches an active product updates that product; otherwise a new one is
  created through :class:`QRProductService`, so the linked inventory Item is
  auto-created exactly like the single-product flow.
* Rows are applied **one transaction at a time**: a bad row is reported and
  skipped without aborting the rest of the file.
* Settings (shelf life / serial prefix) may be supplied either as an id or as
  the setting's human ``value`` — both are exported for round-tripping.
"""

from __future__ import annotations

import csv
import io
import logging
from datetime import UTC, datetime
from decimal import Decimal, InvalidOperation
from uuid import UUID

from fastapi import HTTPException
from sqlalchemy import func, or_
from sqlalchemy.exc import IntegrityError, SQLAlchemyError
from sqlalchemy.orm import Session, joinedload

from app.models.item import Item
from app.models.qr_product import QRProduct
from app.models.qr_product_setting import QRProductSetting
from app.schemas.qr_product import (
    QRProductCreate,
    QRProductPackagingDetails,
    QRProductUpdate,
)
from app.services.qr_product_service import (
    QRProductService,
    find_active_product_by_sku,
    normalize_sku,
)

logger = logging.getLogger(__name__)

#: Columns a file may carry. Order defines the template/export layout.
PRODUCT_COLUMNS: list[str] = [
    "name",
    "sku",
    "generic_name",
    "gtin",
    "industry",
    "email",
    "phone_number",
    "activation_method",
    "sr_number_type",
    "qr_type",
    "warranty_period_months",
    "redirect_to_client",
    "is_active",
    "brand_id",
    "shelf_life_setting_id",
    "shelf_life_setting_value",
    "serial_prefix_setting_id",
    "serial_prefix_setting_value",
    # Inner carton / base packaging unit (IC)
    "unit_name",
    "conversion_factor",
    "items_per_master_pack",
    "length_mm",
    "width_mm",
    "height_mm",
    "weight_grams",
    # Master carton (MC)
    "master_pack_unit_name",
    "master_pack_length_mm",
    "master_pack_width_mm",
    "master_pack_height_mm",
    "master_pack_weight_grams",
    "master_pack_fill_factor",
    "master_pack_void_fill_pct",
    "master_pack_wall_thickness_mm",
    # Linked inventory item (mirrors the item-creation form)
    "description",
    "uom",
    "item_group_id",
    "item_group_name",
    "maintain_stock",
    "valuation_method",
    "standard_rate",
    "valuation_rate",
    "min_order_qty",
    "max_order_qty",
    "reorder_level",
    "reorder_qty",
    "weight_per_unit",
    "weight_uom",
    "barcode",
    "image_url",
    "item_status",
]

#: Extra, read-only columns included in exports only.
EXPORT_ONLY_COLUMNS: list[str] = [
    "id",
    "linked_item_id",
    "linked_item_code",
    "serial_prefix",
    "created_at",
    "updated_at",
]

#: Inner carton / base packaging unit (IC) columns.
_IC_PACKAGING_COLUMNS = (
    "unit_name",
    "conversion_factor",
    "items_per_master_pack",
    "length_mm",
    "width_mm",
    "height_mm",
    "weight_grams",
)

#: Master carton (MC) columns.
_MC_PACKAGING_COLUMNS = (
    "master_pack_unit_name",
    "master_pack_length_mm",
    "master_pack_width_mm",
    "master_pack_height_mm",
    "master_pack_weight_grams",
    "master_pack_fill_factor",
    "master_pack_void_fill_pct",
    "master_pack_wall_thickness_mm",
)

_PACKAGING_COLUMNS = _IC_PACKAGING_COLUMNS + _MC_PACKAGING_COLUMNS

#: Columns forwarded to the linked inventory item.
_ITEM_TEXT_FIELDS = ("description", "uom", "weight_uom", "barcode", "image_url")
_ITEM_DECIMAL_FIELDS = ("standard_rate", "valuation_rate", "weight_per_unit")
_ITEM_INT_FIELDS = ("min_order_qty", "max_order_qty", "reorder_level", "reorder_qty")
_ITEM_BOOL_FIELDS = ("maintain_stock",)

_ITEM_STATUSES = {"draft", "pending_approval", "active", "inactive", "discontinued"}
_VALUATION_METHODS = {"fifo", "lifo", "moving_average", "standard"}

_BOOL_TRUE = {"1", "true", "yes", "y", "t"}

CSV_MEDIA_TYPE = "text/csv"
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

MAX_IMPORT_ROWS = 5000


# ── primitive coercion ────────────────────────────────────────────────────────


def _as_text(value) -> str | None:
    if value is None:
        return None
    if isinstance(value, str):
        cleaned = value.strip()
        return cleaned or None
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    cleaned = str(value).strip()
    return cleaned or None


def _to_bool(value, *, field: str, default: bool = False) -> bool:
    text = _as_text(value)
    if text is None:
        return default
    lowered = text.lower()
    if lowered in _BOOL_TRUE:
        return True
    if lowered in {"0", "false", "no", "n", "f"}:
        return False
    raise ValueError(f"{field} must be a boolean (true/false)")


def _to_int(value, *, field: str) -> int | None:
    text = _as_text(value)
    if text is None:
        return None
    try:
        number = Decimal(text)
    except InvalidOperation as exc:
        raise ValueError(f"{field} must be a whole number") from exc
    # ``int(Decimal("1.9"))`` silently truncates to 1 — reject fractions
    # instead of accepting a value the error message says must be whole.
    if number != number.to_integral_value():
        raise ValueError(f"{field} must be a whole number")
    return int(number)


def _to_decimal(value, *, field: str) -> Decimal | None:
    text = _as_text(value)
    if text is None:
        return None
    try:
        return Decimal(text)
    except InvalidOperation as exc:
        raise ValueError(f"{field} must be a number") from exc


def _maybe_uuid(value) -> UUID | None:
    text = _as_text(value)
    if text is None:
        return None
    try:
        return UUID(text)
    except (ValueError, AttributeError):
        return None


def _enum_value(value):
    """Return the plain string behind a SQLAlchemy enum column."""
    return getattr(value, "value", value)


def _packaging_from_item(item: Item | None) -> tuple[dict, dict]:
    """Split the linked item's packaging rows into IC and MC field dicts.

    IC is the base ("Each" level) row; MC is the active non-base row. Returns
    empty dicts when the item has no packaging, so the caller can fall back to
    the packaging payload stored on the product.
    """
    if item is None:
        return {}, {}

    units = getattr(item, "packaging_units", None) or []
    base = next((u for u in units if u.is_base_unit and u.is_active), None)
    # A flat row can only describe one master carton. Pick the outermost active
    # outer unit deterministically (largest conversion factor, then name) rather
    # than relying on the relationship's arbitrary row order.
    carton = max(
        (u for u in units if not u.is_base_unit and u.is_active),
        key=lambda u: (u.conversion_factor or 0, u.unit_name or ""),
        default=None,
    )

    ic: dict = {}
    if base is not None:
        ic = {
            "unit_name": base.unit_name,
            "conversion_factor": base.conversion_factor,
            "items_per_master_pack": base.items_per_master_pack,
            "length_mm": base.length_mm,
            "width_mm": base.width_mm,
            "height_mm": base.height_mm,
            "weight_grams": base.weight_grams,
        }

    mc: dict = {}
    if carton is not None:
        mc = {
            "master_pack_unit_name": carton.unit_name,
            "master_pack_length_mm": carton.length_mm,
            "master_pack_width_mm": carton.width_mm,
            "master_pack_height_mm": carton.height_mm,
            "master_pack_weight_grams": carton.weight_grams,
            "master_pack_fill_factor": carton.master_pack_fill_factor,
            "master_pack_void_fill_pct": carton.master_pack_void_fill_pct,
            "master_pack_wall_thickness_mm": carton.master_pack_wall_thickness_mm,
        }

    return ic, mc


def parse_row(raw: dict, row_number: int) -> dict:
    """Normalise one raw file row into a validated payload.

    Returns only the columns that were actually provided (plus the always-
    required ``name``) so an upsert never blanks out untouched fields.
    Raises ``ValueError`` with a human message for an invalid row.
    """
    payload: dict = {"row": row_number}

    name = _as_text(raw.get("name"))
    if not name:
        raise ValueError("name is required")
    payload["name"] = name

    sku = normalize_sku(_as_text(raw.get("sku")))
    if sku is not None:
        payload["sku"] = sku

    _parse_text_fields(raw, payload)
    _parse_activation(raw, payload)
    _parse_warranty(raw, payload)
    _parse_bools(raw, payload)
    _parse_brand(raw, payload)
    _parse_settings(raw, payload)

    packaging = _parse_packaging(raw)
    if packaging:
        payload["packaging"] = packaging

    item_fields = _parse_item_fields(raw)
    if item_fields:
        payload["item_fields"] = item_fields

    return payload


_TEXT_FIELDS = ("generic_name", "gtin", "industry", "email", "phone_number")


def _parse_text_fields(raw: dict, payload: dict) -> None:
    for field in _TEXT_FIELDS:
        text = _as_text(raw.get(field))
        if text is not None:
            payload[field] = text


def _parse_activation(raw: dict, payload: dict) -> None:
    activation = _as_text(raw.get("activation_method"))
    if activation is not None:
        activation = activation.lower()
        if activation not in {"pre", "post"}:
            raise ValueError("activation_method must be 'pre' or 'post'")
        payload["activation_method"] = activation

    for field in ("sr_number_type", "qr_type"):
        text = _as_text(raw.get(field))
        if text is not None:
            payload[field] = text


def _parse_warranty(raw: dict, payload: dict) -> None:
    warranty = _to_int(
        raw.get("warranty_period_months"), field="warranty_period_months"
    )
    if warranty is not None:
        if warranty < 0:
            raise ValueError("warranty_period_months cannot be negative")
        payload["warranty_period_months"] = warranty


def _parse_bools(raw: dict, payload: dict) -> None:
    for field in ("redirect_to_client", "is_active"):
        if _as_text(raw.get(field)) is not None:
            payload[field] = _to_bool(raw.get(field), field=field)


def _parse_brand(raw: dict, payload: dict) -> None:
    if _as_text(raw.get("brand_id")) is None:
        return
    brand_id = _maybe_uuid(raw.get("brand_id"))
    if brand_id is None:
        raise ValueError("brand_id must be a valid UUID")
    payload["brand_id"] = brand_id


def _parse_settings(raw: dict, payload: dict) -> None:
    """Accept each setting as an id or as its human ``value``."""
    for prefix in ("shelf_life_setting", "serial_prefix_setting"):
        setting_id = _maybe_uuid(raw.get(f"{prefix}_id"))
        setting_value = _as_text(raw.get(f"{prefix}_value"))
        raw_id_text = _as_text(raw.get(f"{prefix}_id"))
        if setting_id is None and setting_value is None and raw_id_text:
            # A code was placed in the *_id column — treat it as a value.
            setting_value = raw_id_text
        if setting_id is not None:
            payload[f"{prefix}_id"] = setting_id
        if setting_value is not None:
            payload[f"{prefix}_value"] = setting_value


def _parse_dims(raw: dict, target: dict, fields: tuple[str, ...]) -> None:
    """Parse non-negative decimal columns into ``target``."""
    for field in fields:
        value = _to_decimal(raw.get(field), field=field)
        if value is not None:
            if value < 0:
                raise ValueError(f"{field} cannot be negative")
            target[field] = value


def _parse_packaging(raw: dict) -> dict:
    packaging: dict = {}

    unit_name = _as_text(raw.get("unit_name"))
    if unit_name is not None:
        packaging["unit_name"] = unit_name

    conversion = _to_decimal(raw.get("conversion_factor"), field="conversion_factor")
    if conversion is not None:
        if conversion <= 0:
            raise ValueError("conversion_factor must be greater than 0")
        packaging["conversion_factor"] = conversion

    per_master = _to_int(
        raw.get("items_per_master_pack"), field="items_per_master_pack"
    )
    if per_master is not None:
        if per_master <= 0:
            raise ValueError("items_per_master_pack must be greater than 0")
        packaging["items_per_master_pack"] = per_master

    _parse_dims(raw, packaging, ("length_mm", "width_mm", "height_mm", "weight_grams"))

    _parse_master_pack(raw, packaging)
    _require_master_pack_size(packaging)

    if not any(key in packaging for key in _PACKAGING_COLUMNS):
        return {}
    packaging.setdefault("unit_name", "Each")
    packaging.setdefault("conversion_factor", Decimal("1"))
    return packaging


def _require_master_pack_size(packaging: dict) -> None:
    """Reject an MC-only row that carries no master-pack size.

    The item service derives the master-pack count from
    ``items_per_master_pack`` (or a ``conversion_factor`` greater than 1) and
    deactivates the carton when neither is supplied — silently discarding the
    master-carton fields. Fail the row instead.
    """
    if not any(key in packaging for key in _MC_PACKAGING_COLUMNS):
        return
    if "items_per_master_pack" in packaging:
        return
    conversion = packaging.get("conversion_factor")
    if conversion is not None and conversion > 1:
        return
    raise ValueError(
        "master_pack_* fields require items_per_master_pack "
        "(or a conversion_factor greater than 1)"
    )


def _parse_master_pack(raw: dict, packaging: dict) -> None:
    """Parse the master carton (MC) columns into ``packaging``."""
    unit_name = _as_text(raw.get("master_pack_unit_name"))
    if unit_name is not None:
        packaging["master_pack_unit_name"] = unit_name

    _parse_dims(
        raw,
        packaging,
        (
            "master_pack_length_mm",
            "master_pack_width_mm",
            "master_pack_height_mm",
            "master_pack_weight_grams",
        ),
    )

    fill = _to_decimal(
        raw.get("master_pack_fill_factor"), field="master_pack_fill_factor"
    )
    if fill is not None:
        if not (Decimal("0") < fill <= Decimal("1")):
            raise ValueError(
                "master_pack_fill_factor must be greater than 0 and at most 1"
            )
        packaging["master_pack_fill_factor"] = fill

    void = _to_decimal(
        raw.get("master_pack_void_fill_pct"), field="master_pack_void_fill_pct"
    )
    if void is not None:
        if not (Decimal("0") <= void <= Decimal("1")):
            raise ValueError("master_pack_void_fill_pct must be between 0 and 1")
        packaging["master_pack_void_fill_pct"] = void

    thickness = _to_decimal(
        raw.get("master_pack_wall_thickness_mm"), field="master_pack_wall_thickness_mm"
    )
    if thickness is not None:
        if thickness < 0:
            raise ValueError("master_pack_wall_thickness_mm cannot be negative")
        packaging["master_pack_wall_thickness_mm"] = thickness


def _parse_item_fields(raw: dict) -> dict:
    """Parse the linked-item columns (uom, group, rates, reorder levels…).

    Values are validated but not resolved — ids are looked up per organization
    later, so a bad id is reported against the row rather than the whole file.
    """
    fields: dict = {}
    _parse_item_text(raw, fields)
    _parse_item_numbers(raw, fields)
    _parse_item_bools(raw, fields)
    _parse_item_group(raw, fields)
    _parse_item_enums(raw, fields)
    return fields


def _parse_item_text(raw: dict, fields: dict) -> None:
    for field in _ITEM_TEXT_FIELDS:
        text = _as_text(raw.get(field))
        if text is not None:
            fields[field] = text


def _parse_item_numbers(raw: dict, fields: dict) -> None:
    for field in _ITEM_DECIMAL_FIELDS:
        value = _to_decimal(raw.get(field), field=field)
        if value is not None:
            if value < 0:
                raise ValueError(f"{field} cannot be negative")
            fields[field] = value

    for field in _ITEM_INT_FIELDS:
        value = _to_int(raw.get(field), field=field)
        if value is None:
            continue
        if value < 0:
            raise ValueError(f"{field} cannot be negative")
        if field == "min_order_qty" and value < 1:
            raise ValueError("min_order_qty must be at least 1")
        fields[field] = value


def _parse_item_bools(raw: dict, fields: dict) -> None:
    for field in _ITEM_BOOL_FIELDS:
        if _as_text(raw.get(field)) is not None:
            fields[field] = _to_bool(raw.get(field), field=field)


def _parse_item_group(raw: dict, fields: dict) -> None:
    group_id_text = _as_text(raw.get("item_group_id"))
    if group_id_text is not None:
        group_id = _maybe_uuid(group_id_text)
        if group_id is None:
            raise ValueError("item_group_id must be a valid UUID")
        fields["item_group_id"] = group_id

    group_name = _as_text(raw.get("item_group_name"))
    if group_name is not None:
        fields["item_group_name"] = group_name


def _parse_item_enums(raw: dict, fields: dict) -> None:
    valuation_method = _as_text(raw.get("valuation_method"))
    if valuation_method is not None:
        valuation_method = valuation_method.lower()
        if valuation_method not in _VALUATION_METHODS:
            raise ValueError(
                "valuation_method must be one of: "
                + ", ".join(sorted(_VALUATION_METHODS))
            )
        fields["valuation_method"] = valuation_method

    item_status = _as_text(raw.get("item_status"))
    if item_status is not None:
        item_status = item_status.lower()
        if item_status not in _ITEM_STATUSES:
            raise ValueError(
                "item_status must be one of: " + ", ".join(sorted(_ITEM_STATUSES))
            )
        fields["status"] = item_status


class QRProductBulkService:
    """Import/export QR products as CSV or XLSX."""

    def __init__(self, db: Session):
        self.db = db
        self.product_service = QRProductService(db)

    # ── Export ────────────────────────────────────────────────────────────

    def export_products(
        self,
        organization_id: UUID,
        *,
        file_format: str = "csv",
        is_active: bool | None = None,
        search: str | None = None,
    ) -> tuple[bytes, str]:
        """Return ``(bytes, filename)`` for every matching product."""
        query = (
            self.db.query(QRProduct)
            .options(
                joinedload(QRProduct.shelf_life_setting),
                joinedload(QRProduct.serial_prefix_setting),
                joinedload(QRProduct.items).joinedload(Item.packaging_units),
                joinedload(QRProduct.items).joinedload(Item.item_group),
            )
            .filter(
                QRProduct.organization_id == organization_id,
                QRProduct.deleted_at.is_(None),
            )
        )
        if is_active is not None:
            query = query.filter(QRProduct.is_active == is_active)
        if search:
            term = f"%{search.strip()}%"
            query = query.filter(
                or_(QRProduct.name.ilike(term), QRProduct.sku.ilike(term))
            )
        products = query.order_by(QRProduct.created_at.asc()).all()

        rows = [self._export_row(product) for product in products]
        columns = PRODUCT_COLUMNS + EXPORT_ONLY_COLUMNS
        return self._render(rows, columns, file_format, "qr_products_export")

    def template(self, file_format: str = "csv") -> tuple[bytes, str]:
        """Return ``(bytes, filename)`` for an import template (headers only)."""
        columns = PRODUCT_COLUMNS
        return self._render([], columns, file_format, "qr_products_import_template")

    @staticmethod
    def _export_row(product: QRProduct) -> dict:
        item = product.items[0] if product.items else None
        stored = (product.extra_data or {}).get("packaging_details") or {}
        ic, mc = _packaging_from_item(item)

        row = {
            "id": str(product.id),
            "name": product.name,
            "sku": product.sku,
            "generic_name": product.generic_name,
            "gtin": product.gtin,
            "industry": product.industry,
            "email": product.email,
            "phone_number": product.phone_number,
            "activation_method": product.activation_method,
            "sr_number_type": product.sr_number_type,
            "qr_type": product.qr_type,
            "warranty_period_months": product.warranty_period_months,
            "redirect_to_client": product.redirect_to_client,
            "is_active": product.is_active,
            "brand_id": str(product.brand_id) if product.brand_id else None,
            "shelf_life_setting_id": str(product.shelf_life_setting_id)
            if product.shelf_life_setting_id
            else None,
            "shelf_life_setting_value": (
                product.shelf_life_setting.value if product.shelf_life_setting else None
            ),
            "serial_prefix_setting_id": str(product.serial_prefix_setting_id)
            if product.serial_prefix_setting_id
            else None,
            "serial_prefix_setting_value": (
                product.serial_prefix_setting.value
                if product.serial_prefix_setting
                else None
            ),
        }

        # Packaging (IC + MC): prefer the live item packaging rows, fall back to
        # the payload stored on the product's extra_data.
        for key in _PACKAGING_COLUMNS:
            if key in ic:
                row[key] = ic[key]
            elif key in mc:
                row[key] = mc[key]
            else:
                row[key] = stored.get(key)

        row.update(
            {
                "description": item.description if item else None,
                "uom": item.uom if item else None,
                "item_group_id": (
                    str(item.item_group_id) if item and item.item_group_id else None
                ),
                "item_group_name": (
                    item.item_group.name if item and item.item_group else None
                ),
                "maintain_stock": item.maintain_stock if item else None,
                "valuation_method": _enum_value(item.valuation_method)
                if item
                else None,
                "standard_rate": item.standard_rate if item else None,
                "valuation_rate": item.valuation_rate if item else None,
                "min_order_qty": item.min_order_qty if item else None,
                "max_order_qty": item.max_order_qty if item else None,
                "reorder_level": item.reorder_level if item else None,
                "reorder_qty": item.reorder_qty if item else None,
                "weight_per_unit": item.weight_per_unit if item else None,
                "weight_uom": item.weight_uom if item else None,
                "barcode": item.barcode if item else None,
                "image_url": item.image_url if item else None,
                "item_status": _enum_value(item.status) if item else None,
                "linked_item_id": str(item.id) if item else None,
                "linked_item_code": item.item_code if item else None,
                "serial_prefix": product.serial_prefix,
                "created_at": product.created_at.isoformat()
                if product.created_at
                else None,
                "updated_at": product.updated_at.isoformat()
                if product.updated_at
                else None,
            }
        )
        return row

    @staticmethod
    def _render(
        rows: list[dict], columns: list[str], file_format: str, stem: str
    ) -> tuple[bytes, str]:
        file_format = (file_format or "csv").lower()
        stamp = datetime.now(UTC).strftime("%Y%m%d_%H%M%S")
        if file_format in {"xlsx", "xlsm"}:
            from openpyxl import Workbook

            wb = Workbook()
            ws = wb.active
            ws.title = "QR Products"
            ws.append(columns)
            for row in rows:
                ws.append([_cell(row.get(col)) for col in columns])
            buffer = io.BytesIO()
            wb.save(buffer)
            return buffer.getvalue(), f"{stem}_{stamp}.xlsx"

        buffer = io.StringIO()
        writer = csv.DictWriter(buffer, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({col: _cell(row.get(col)) for col in columns})
        return buffer.getvalue().encode("utf-8-sig"), f"{stem}_{stamp}.csv"

    # ── Import ────────────────────────────────────────────────────────────

    def import_products(
        self,
        organization_id: UUID,
        user_id: UUID,
        file_bytes: bytes,
        filename: str | None,
    ) -> dict:
        """Upsert products from a CSV/XLSX file and return a per-row summary."""
        rows = self._read_rows(file_bytes, filename)
        if len(rows) > MAX_IMPORT_ROWS:
            raise ValueError(
                f"File has {len(rows)} rows; the maximum is {MAX_IMPORT_ROWS}"
            )

        created = updated = 0
        errors: list[dict] = []
        for index, raw in enumerate(rows, start=2):  # row 1 is the header
            try:
                payload = parse_row(raw, index)
                was_created = self._upsert_row(organization_id, user_id, payload)
            except (HTTPException, ValueError, IntegrityError) as exc:
                # Expected, row-scoped failures: record and keep going.
                self.db.rollback()
                message = getattr(exc, "detail", None) or str(exc)
                errors.append(
                    {
                        "row": index,
                        "sku": _as_text(raw.get("sku")),
                        "message": str(message),
                    }
                )
                continue
            except SQLAlchemyError:
                # A database outage / programming error is not a row problem —
                # surface it instead of reporting a misleading partial success.
                self.db.rollback()
                raise
            created += was_created
            updated += not was_created

        return {
            "total_rows": len(rows),
            "created": created,
            "updated": updated,
            "failed": len(errors),
            "errors": errors,
        }

    def _upsert_row(self, organization_id: UUID, user_id: UUID, payload: dict) -> bool:
        """Create or update one product. Returns True when a row was created."""
        packaging = payload.get("packaging")
        item_fields = payload.get("item_fields")
        if payload.get("brand_id") is not None:
            self._assert_brand(organization_id, payload["brand_id"])

        # Resolve item references before any write: the product create/update
        # below commits internally, so an unknown item group must fail the row
        # while nothing has been persisted yet.
        if item_fields:
            item_fields = self._prepare_item_fields(organization_id, item_fields)

        existing = find_active_product_by_sku(
            self.db, organization_id, payload.get("sku")
        )
        packaging_details = (
            QRProductPackagingDetails(**packaging) if packaging else None
        )

        # Settings are only resolved when the row supplies them, so a partial
        # update (e.g. name + sku) preserves the existing shelf-life /
        # serial-prefix instead of failing on a missing setting.
        if existing is not None:
            update = QRProductUpdate(
                **self._update_kwargs(payload),
                **self._resolved_settings(organization_id, payload, required=False),
                packaging_details=packaging_details,
            )
            product = self.product_service.update_product(
                existing.id, update, organization_id, user_id
            )
            created = False
        else:
            create = QRProductCreate(
                **self._create_kwargs(payload),
                **self._resolved_settings(organization_id, payload, required=True),
                packaging_details=packaging_details,
            )
            product = self.product_service.create_product(
                create, organization_id, user_id
            )
            created = True

        if item_fields:
            self._apply_item_fields(product, item_fields, organization_id, user_id)
        return created

    def _apply_item_fields(
        self,
        product: QRProduct,
        fields: dict,
        organization_id: UUID,
        user_id: UUID,
    ) -> None:
        """Write the row's item-level columns onto the linked inventory item."""
        from app.schemas.item import ItemUpdate
        from app.services.item_service import ItemService

        item = (
            self.db.query(Item)
            .filter(Item.qr_product_id == product.id, Item.deleted_at.is_(None))
            .first()
        )
        if item is None:
            logger.warning(
                "QR product %s has no linked item — skipping item fields", product.id
            )
            return

        ItemService(self.db).update_item(
            item.id, ItemUpdate(**fields), organization_id, user_id
        )

    def _prepare_item_fields(self, organization_id: UUID, fields: dict) -> dict:
        """Resolve the row's item group to an id before anything is written."""
        resolved = dict(fields)
        group_ref = resolved.pop("item_group_name", None)
        if "item_group_id" in resolved:
            self._assert_item_group(organization_id, resolved["item_group_id"])
        elif group_ref is not None:
            resolved["item_group_id"] = self._resolve_item_group(
                organization_id, group_ref
            )
        return resolved

    def _resolve_item_group(self, organization_id: UUID, reference: str) -> UUID:
        """Resolve an item group by (case-insensitive) name or code."""
        from app.models.item_group import ItemGroup

        group = (
            self.db.query(ItemGroup)
            .filter(
                ItemGroup.organization_id == organization_id,
                ItemGroup.deleted_at.is_(None),
                or_(
                    func.lower(ItemGroup.name) == reference.lower(),
                    func.lower(ItemGroup.code) == reference.lower(),
                ),
            )
            .first()
        )
        if group is None:
            raise ValueError(f"item group '{reference}' was not found")
        return group.id

    def _assert_item_group(self, organization_id: UUID, group_id: UUID) -> None:
        from app.models.item_group import ItemGroup

        exists = (
            self.db.query(ItemGroup.id)
            .filter(
                ItemGroup.id == group_id,
                ItemGroup.organization_id == organization_id,
                ItemGroup.deleted_at.is_(None),
            )
            .first()
        )
        if exists is None:
            raise ValueError(f"item group '{group_id}' was not found")

    def _resolved_settings(
        self, organization_id: UUID, payload: dict, *, required: bool
    ) -> dict:
        """Resolve whichever settings the row provided (by id or by value)."""
        settings: dict = {}
        for prefix, setting_type in (
            ("shelf_life_setting", "shelf_life"),
            ("serial_prefix_setting", "serial_prefix"),
        ):
            provided = f"{prefix}_id" in payload or f"{prefix}_value" in payload
            if not provided:
                if required:
                    raise ValueError(f"{prefix}_id or {prefix}_value is required")
                continue
            settings[f"{prefix}_id"] = self._resolve_setting(
                organization_id, payload, prefix, setting_type
            )
        return settings

    @staticmethod
    def _create_kwargs(payload: dict) -> dict:
        allowed = (
            "name",
            "sku",
            "generic_name",
            "gtin",
            "brand_id",
            "industry",
            "email",
            "phone_number",
            "activation_method",
            "sr_number_type",
            "qr_type",
            "warranty_period_months",
            "redirect_to_client",
        )
        return {key: payload[key] for key in allowed if key in payload}

    @staticmethod
    def _update_kwargs(payload: dict) -> dict:
        # ``sku`` is the upsert key and ``brand_id`` is immutable after create.
        allowed = (
            "name",
            "generic_name",
            "gtin",
            "industry",
            "email",
            "phone_number",
            "activation_method",
            "sr_number_type",
            "qr_type",
            "warranty_period_months",
            "redirect_to_client",
            "is_active",
        )
        return {key: payload[key] for key in allowed if key in payload}

    def _resolve_setting(
        self,
        organization_id: UUID,
        payload: dict,
        prefix: str,
        setting_type: str,
    ) -> UUID:
        """Resolve a setting to its id from either the id or the human value."""
        setting_id = payload.get(f"{prefix}_id")
        setting_value = payload.get(f"{prefix}_value")
        if setting_id is None and setting_value is None:
            raise ValueError(f"{prefix}_id or {prefix}_value is required")

        setting = None
        if setting_id is not None:
            setting = (
                self.db.query(QRProductSetting)
                .filter(
                    QRProductSetting.id == setting_id,
                    QRProductSetting.organization_id == organization_id,
                    QRProductSetting.deleted_at.is_(None),
                )
                .first()
            )
        elif setting_value is not None:
            setting = (
                self.db.query(QRProductSetting)
                .filter(
                    QRProductSetting.organization_id == organization_id,
                    QRProductSetting.setting_type == setting_type,
                    func.lower(QRProductSetting.value) == setting_value.lower(),
                    QRProductSetting.deleted_at.is_(None),
                )
                .first()
            )

        if setting is None:
            label = setting_value or setting_id
            raise ValueError(f"{setting_type} setting '{label}' was not found")
        if setting.setting_type != setting_type:
            raise ValueError(f"Setting '{label}' is not a {setting_type} setting")
        if not setting.is_active:
            raise ValueError(f"{setting_type} setting '{label}' is inactive")
        return setting.id

    def _assert_brand(self, organization_id: UUID, brand_id: UUID) -> None:
        from app.repositories.brand_repository import BrandRepository

        if BrandRepository(self.db).get_by_id(brand_id, organization_id) is None:
            raise ValueError(f"Brand '{brand_id}' was not found")

    # ── File parsing ──────────────────────────────────────────────────────

    def _read_rows(self, file_bytes: bytes, filename: str | None) -> list[dict]:
        if not file_bytes:
            raise ValueError("Uploaded file is empty")
        ext = (
            filename.rsplit(".", 1)[-1].lower()
            if filename and "." in filename
            else "csv"
        )
        if ext in {"xlsx", "xlsm", "xltx"}:
            rows = self._read_xlsx(file_bytes)
        elif ext in {"csv", "txt"}:
            rows = self._read_csv(file_bytes)
        else:
            raise ValueError(f"Unsupported file type '.{ext}'. Supported: csv, xlsx")
        return [row for row in rows if self._has_values(row)]

    @staticmethod
    def _has_values(row: dict) -> bool:
        return any(_as_text(value) is not None for value in row.values())

    @staticmethod
    def _read_csv(file_bytes: bytes) -> list[dict]:
        text = file_bytes.decode("utf-8-sig", errors="replace")
        reader = csv.DictReader(io.StringIO(text))
        rows = []
        for raw in reader:
            rows.append(
                {
                    (key.strip() if isinstance(key, str) else key): value
                    for key, value in raw.items()
                    if key is not None
                }
            )
        return rows

    @staticmethod
    def _read_xlsx(file_bytes: bytes) -> list[dict]:
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        header = next(iterator, None)
        if not header:
            return []
        columns = [(_as_text(cell) or "") for cell in header]
        rows = []
        for values in iterator:
            row = {
                columns[idx]: values[idx] for idx in range(len(columns)) if columns[idx]
            }
            rows.append(row)
        return rows


#: Leading characters spreadsheet software may interpret as a formula.
_FORMULA_TRIGGERS = ("=", "+", "-", "@", "\t", "\r")


def _cell(value):
    """Render a value for CSV/XLSX output, neutralising formula injection.

    A user-controlled string such as ``=HYPERLINK(...)`` or ``+cmd|...`` would
    be executed as a formula when the file is opened in Excel/Sheets; prefixing
    it with an apostrophe keeps it inert text.
    """
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, str) and value[:1] in _FORMULA_TRIGGERS:
        return "'" + value
    return value
